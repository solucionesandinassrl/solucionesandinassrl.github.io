"""Generador de firmas de correo de Soluciones Andinas.

Lee un Excel con las personas, valida los datos y genera una firma HTML y una de texto por persona,
mas un mapeo para DevOps y un reporte con la vista previa de todas las firmas.

Uso:
    python generar_firmas.py personas.xlsx
    python generar_firmas.py personas.xlsx --salida D:/firmas --plantilla plantillas/firma.html

Ver README.md.
"""
import argparse, csv, html, json, os, re, shutil, sys, unicodedata, urllib.parse
from datetime import datetime

try:
    import openpyxl
except ImportError:
    sys.exit("Falta la libreria openpyxl. Instalala con:  python -m pip install openpyxl")

AQUI = os.path.dirname(os.path.abspath(__file__))
FIRMAS_URL = "https://solucionesandinassrl.github.io/firmas"
FIRMAS_LOCAL = os.path.join(os.path.dirname(AQUI), "firmas")  # para la vista previa del reporte
ASSETS_URL = FIRMAS_URL + "/assets"
FOTOS_LOCAL = os.path.join(FIRMAS_LOCAL, "fotos")  # fotos circulares: <usuario>.png (ver preparar_fotos.py)
FIJOS = {"empresa": "Soluciones Andinas", "web": "https://solucionesandinas.com.ar",
         "web_vis": "solucionesandinas.com.ar", "assets": ASSETS_URL}
COLUMNAS = {"email": "email", "nombre": "nombre", "apellido": "apellido", "cargo": "cargo",
            "area": "area", "celular": "celular", "celular_firma": "celular en firma",
            "linkedin": "linkedin", "incluir": "incluir"}
SI = ("si", "s", "1", "true", "x")
OPCIONALES = ("linkedin",)


def sin_tilde(s):
    return unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()


def normalizar_celular(valor):
    """Devuelve (e164, visible) o lanza ValueError. Regla del 9 para celulares argentinos."""
    d = re.sub(r"\D", "", str(valor or ""))
    if d.startswith("549"):
        d = d[3:]
    elif d.startswith("54"):
        d = d[2:]
    d = d.lstrip("0")
    if len(d) != 10:
        raise ValueError(f"celular '{valor}': se esperaban 10 digitos (caracteristica + numero, sin 0 ni 15)")
    e164 = "+549" + d
    if d.startswith("11"):  # AMBA: caracteristica de 2 digitos
        vis = f"+54 9 11 {d[2:6]}-{d[6:]}"
    else:  # caracteristica de 3 digitos (ej. 261 Mendoza)
        vis = f"+54 9 {d[:3]} {d[3:6]}-{d[6:]}"
    return e164, vis


def normalizar_linkedin(valor):
    v = str(valor or "").strip()
    if not v:
        return ""
    if not re.match(r"^https?://", v, re.I):
        v = "https://" + v
    p = urllib.parse.urlsplit(v)
    if "linkedin.com" not in p.netloc.lower():
        raise ValueError(f"LinkedIn '{valor}': no es una URL de linkedin.com")
    ruta = urllib.parse.quote(urllib.parse.unquote(p.path), safe="/-_.~")
    return urllib.parse.urlunsplit(("https", p.netloc.lower(), ruta, "", ""))


def leer_excel(ruta):
    wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    ws = wb["Personas"] if "Personas" in wb.sheetnames else wb.worksheets[0]
    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        sys.exit("El Excel esta vacio.")
    enc = [sin_tilde(str(c or "")).strip().lower() for c in filas[0]]
    idx = {}
    for clave, nombre in COLUMNAS.items():
        if nombre in enc:
            idx[clave] = enc.index(nombre)
    faltan = [c for c in ("email", "nombre", "apellido", "cargo", "celular") if c not in idx]
    if faltan:
        sys.exit(f"Faltan columnas en el Excel: {', '.join(faltan)}. Usa plantilla-personas.xlsx")
    personas = []
    for n, fila in enumerate(filas[1:], start=2):
        if not any(c not in (None, "") for c in fila):
            continue
        reg = {k: ("" if fila[i] is None else str(fila[i]).strip()) for k, i in idx.items() if i < len(fila)}
        reg["_fila"] = n
        personas.append(reg)
    return personas


def validar(personas):
    ok, errores, avisos, vistos = [], [], [], {}
    for p in personas:
        motivos = []
        if sin_tilde(p.get("incluir", "si")).lower() in ("no", "n", "0", "false"):
            continue
        email = p.get("email", "").lower()
        if not re.match(r"^[a-z0-9._%+-]+@[a-z0-9.-]+\.[a-z]{2,}$", email):
            motivos.append(f"email invalido '{p.get('email', '')}'")
        elif email in vistos:
            motivos.append(f"email repetido (tambien en la fila {vistos[email]})")
        else:
            vistos[email] = p["_fila"]
            if not email.endswith("@sandinas.com.ar"):
                avisos.append((p["_fila"], email, "el email no es @sandinas.com.ar"))
        for campo in ("nombre", "apellido", "cargo"):
            if not p.get(campo):
                motivos.append(f"falta {campo}")
        mostrar_cel = sin_tilde(p.get("celular_firma", "")).strip().lower() in SI
        try:
            p["celular_e164"], p["celular"] = normalizar_celular(p.get("celular"))
            # celular_firma: el celular, solo si la persona lo quiere en la firma (columna "Celular en firma" = Si)
            p["celular_firma"] = p["celular"] if mostrar_cel else ""
        except ValueError as e:
            motivos.append(str(e))
        try:
            p["linkedin"] = normalizar_linkedin(p.get("linkedin"))
        except ValueError as e:
            motivos.append(str(e))
        if motivos:
            errores.append((p["_fila"], p.get("email", ""), "; ".join(motivos)))
        else:
            p["email"] = email
            p["usuario"] = re.sub(r"[^a-z0-9._-]", "", email.split("@")[0])
            # foto opcional: si existe firmas/fotos/<usuario>.png se usa su URL publica
            p["foto"] = f"{FIRMAS_URL}/fotos/{p['usuario']}.png" if os.path.exists(
                os.path.join(FOTOS_LOCAL, p["usuario"] + ".png")) else ""
            ok.append(p)
    return ok, errores, avisos


def completar(plantilla, datos, es_html):
    """Reemplaza {{campo}} y quita los bloques <!--SI:campo-->...<!--FIN:campo--> si el campo esta vacio."""
    def bloque(m):
        return m.group(2) if datos.get(m.group(1)) else ""
    t = re.sub(r"<!--SI:(\w+)-->(.*?)<!--FIN:\1-->", bloque, plantilla, flags=re.S)
    # version texto: si el bloque debe ocupar su propia linea, el salto de linea va DENTRO del bloque
    t = re.sub(r"\[\[SI:(\w+)\]\](.*?)\[\[FIN:\1\]\]", bloque, t, flags=re.S)
    def valor(m):
        v = str(datos.get(m.group(1), ""))
        return html.escape(v, quote=True) if es_html else v
    t = re.sub(r"\{\{(\w+)\}\}", valor, t)
    pendientes = re.findall(r"\{\{(\w+)\}\}", t)
    if pendientes:
        raise ValueError(f"la plantilla usa campos desconocidos: {', '.join(set(pendientes))}")
    return t


def reporte(destino, generadas, errores, avisos, plantilla_html):
    hoy = datetime.now().strftime("%d/%m/%Y %H:%M")
    local = os.path.isdir(FIRMAS_LOCAL)
    def vista(h):  # en la vista previa usa los archivos locales si existen (la URL publica puede no estar subida)
        return h.replace(FIRMAS_URL, "file:///" + FIRMAS_LOCAL.replace("\\", "/")) if local else h
    filas_err = "".join(f"<tr><td>{f}</td><td>{html.escape(e)}</td><td>{html.escape(m)}</td></tr>" for f, e, m in errores)
    filas_av = "".join(f"<tr><td>{f}</td><td>{html.escape(e)}</td><td>{html.escape(m)}</td></tr>" for f, e, m in avisos)
    tarjetas = "".join(
        f'<section class="f"><div class="cab"><b>{html.escape(p["nombre"])} {html.escape(p["apellido"])}</b>'
        f'<span>{html.escape(p["archivo_html"])}</span><button type="button" data-id="s{i}">Copiar firma</button></div>'
        f'<div class="firma" id="s{i}">{vista(p["html"])}</div></section>'
        for i, p in enumerate(generadas))
    doc = f"""<!doctype html><html lang="es"><head><meta charset="utf-8"><title>Reporte de firmas</title>
<style>
body{{margin:0;background:#F3F5F7;font:14px "Segoe UI",Arial,sans-serif;color:#0C3B51}}
.w{{max-width:980px;margin:0 auto;padding:24px 16px;display:flex;flex-direction:column;gap:16px}}
h1{{margin:0;font-size:22px}} .sub{{color:#656666;margin:4px 0 0}}
.kpi{{display:flex;gap:12px;flex-wrap:wrap}} .kpi div{{background:#fff;border:1px solid #E6EAED;border-radius:10px;padding:10px 16px}}
.kpi b{{font-size:22px;display:block}} .ok b{{color:#1F9D63}} .ko b{{color:#D2463C}} .av b{{color:#B7791F}}
.t{{border-collapse:collapse;width:100%;background:#fff;border:1px solid #E6EAED}} .t td,.t th{{padding:7px 10px;border-bottom:1px solid #E6EAED;text-align:left}}
h2{{font-size:15px;margin:8px 0 0}}
.f{{background:#fff;border:1px solid #E6EAED;border-radius:10px;overflow:hidden}}
.cab{{display:flex;gap:12px;align-items:center;padding:8px 14px;background:#F7F9FA;border-bottom:1px solid #E6EAED}}
.cab span{{color:#656666;font-size:12px;margin-right:auto}}
.cab button{{font:600 12px "Segoe UI",Arial;background:#1798D5;color:#fff;border:0;border-radius:8px;padding:6px 12px;cursor:pointer}}
.firma{{padding:18px}}
</style></head><body><div class="w">
<div><h1>Reporte de firmas</h1><p class="sub">Generado el {hoy} · plantilla {html.escape(plantilla_html)}</p></div>
<div class="kpi"><div class="ok"><b>{len(generadas)}</b>firmas generadas</div><div class="ko"><b>{len(errores)}</b>filas con error</div><div class="av"><b>{len(avisos)}</b>avisos</div></div>
{"<h2>Filas con error (no se generaron)</h2><table class='t'><tr><th>Fila</th><th>Email</th><th>Motivo</th></tr>" + filas_err + "</table>" if errores else ""}
{"<h2>Avisos (se generaron igual)</h2><table class='t'><tr><th>Fila</th><th>Email</th><th>Aviso</th></tr>" + filas_av + "</table>" if avisos else ""}
<h2>Vista previa</h2>{tarjetas}
</div>
<script>
var LOCAL = {json.dumps("file:///" + FIRMAS_LOCAL.replace(chr(92), "/"))}, PUBLICA = {json.dumps(FIRMAS_URL)};
document.addEventListener("click", function(e){{
  var b = e.target.closest("button[data-id]"); if(!b) return;
  var el = document.getElementById(b.dataset.id), r = document.createRange(); r.selectNodeContents(el);
  // la vista previa puede mostrar el logo desde la PC; lo que se copia siempre usa la URL publica
  var imgs = el.querySelectorAll("img"), orig = [];
  imgs.forEach(function(i){{ orig.push(i.getAttribute("src")); i.setAttribute("src", i.getAttribute("src").replace(LOCAL, PUBLICA)); }});
  var s = window.getSelection(); s.removeAllRanges(); s.addRange(r);
  var ok = false; try {{ ok = document.execCommand("copy"); }} catch(x) {{}}
  s.removeAllRanges(); imgs.forEach(function(i, k){{ i.setAttribute("src", orig[k]); }});
  b.textContent = ok ? "Copiada" : "Seleccioná y copiá (Ctrl+C)";
  setTimeout(function(){{ b.textContent = "Copiar firma"; }}, 2000);
}});
</script></body></html>"""
    open(os.path.join(destino, "reporte.html"), "w", encoding="utf-8").write(doc)


def main():
    ap = argparse.ArgumentParser(description="Genera las firmas de correo desde un Excel.")
    ap.add_argument("excel", help="Excel con las personas (formato de plantilla-personas.xlsx)")
    ap.add_argument("--salida", default=os.path.join(AQUI, "salida"), help="carpeta donde dejar las firmas")
    ap.add_argument("--plantilla", default=os.path.join(AQUI, "plantillas", "firma.html"),
                    help="plantilla HTML (la de texto se busca con el mismo nombre y extension .txt)")
    a = ap.parse_args()

    t_html = open(a.plantilla, encoding="utf-8").read()
    ruta_txt = os.path.splitext(a.plantilla)[0] + ".txt"
    t_txt = open(ruta_txt, encoding="utf-8").read() if os.path.exists(ruta_txt) else None

    ok, errores, avisos = validar(leer_excel(a.excel))
    destino = os.path.join(a.salida, "firmas_" + datetime.now().strftime("%Y%m%d"))
    shutil.rmtree(destino, ignore_errors=True)
    os.makedirs(destino)

    generadas, mapeo = [], []
    ahora = datetime.now().isoformat(timespec="seconds")
    for p in ok:
        datos = {**FIJOS, **p}
        try:
            h = completar(t_html, datos, True)
            t = completar(t_txt, datos, False) if t_txt else None
        except ValueError as e:
            errores.append((p["_fila"], p["email"], str(e))); continue
        p["archivo_html"] = f"firma_{p['usuario']}.html"
        open(os.path.join(destino, p["archivo_html"]), "w", encoding="utf-8").write(h)
        archivo_txt = ""
        if t is not None:
            archivo_txt = f"firma_{p['usuario']}.txt"
            open(os.path.join(destino, archivo_txt), "w", encoding="utf-8").write(t)
        p["html"] = h
        generadas.append(p)
        mapeo.append([p["email"], p["nombre"], p["apellido"], p["archivo_html"], archivo_txt,
                      os.path.basename(a.plantilla), ahora])

    with open(os.path.join(destino, "mapeo.csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["email", "nombre", "apellido", "archivo_html", "archivo_txt", "plantilla", "generada_en"])
        w.writerows(mapeo)
    reporte(destino, generadas, errores, avisos, os.path.basename(a.plantilla))

    print(f"Firmas generadas: {len(generadas)}  |  Filas con error: {len(errores)}  |  Avisos: {len(avisos)}")
    for f, e, m in errores:
        print(f"  ERROR fila {f} ({e}): {m}")
    for f, e, m in avisos:
        print(f"  aviso fila {f} ({e}): {m}")
    print(f"Salida: {destino}")
    print(f"Revisa el reporte: {os.path.join(destino, 'reporte.html')}")
    return 1 if errores else 0


if __name__ == "__main__":
    sys.exit(main())
